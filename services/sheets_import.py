"""Import the old Google Sheet (downloaded as Excel: File → Download → Microsoft Excel) into one organization:
«الاقسام» → departments, «النماذج» + «الردود» → clearance forms and responses, «العهد» → custody.
Employees met in the forms are added to the roster (inactive — they have left). Running it again skips what
is already imported (departments by code, forms by their FORM-… id, custody by CUST-… id)."""
import io
import re
from datetime import date, datetime

from extensions import db
from models import Department, ClearanceForm, ClearanceResponse, Employee
from models.core import utcnow
from utils.tabular import norm, to_records

DEPT_FIELDS = [('legacy_id', ['id', 'المعرف', 'معرف']), ('code', ['code', 'الكود', 'كود']), ('name', ['name', 'الاسم', 'القسم']),
               ('icon', ['icon', 'الايقونه', 'أيقونة']), ('manager', ['manager', 'المسؤول', 'المدير']),
               ('email', ['email', 'البريد']), ('type', ['type', 'النوع'])]
FORM_FIELDS = [('form_id', ['formId', 'id']), ('ref', ['formCode', 'code', 'الرقم المرجعي']), ('emp_code', ['empCode']),
               ('name', ['employeeName', 'الاسم', 'اسم الموظف']), ('number', ['employeeId', 'الرقم الوظيفي']),
               ('email', ['employeeEmail', 'البريد']), ('department', ['department', 'القسم']),
               ('resign', ['resignDate', 'تاريخ الاستقالة']), ('created', ['createdAt', 'تاريخ الإنشاء']),
               ('status', ['status', 'الحالة']), ('hr_note', ['hrNote', 'ملاحظة'])]
RESP_FIELDS = [('form_id', ['formId']), ('dept_id', ['deptId']), ('status', ['status', 'الحالة']),
               ('note', ['note', 'ملاحظات']), ('at', ['respondedAt']), ('by', ['respondedBy'])]

FORM_STATUS = {norm('قيد التنفيذ'): 'in_progress', norm('مكتمل'): 'completed', norm('مقفل'): 'closed',
               norm('ملغي'): 'cancelled'}
RESP_STATUS = {norm('قيد الانتظار'): 'pending', norm('خالية الذمة'): 'cleared', norm('معلق'): 'held'}
DEPT_TYPE = {'hr': 'HR', 'finance': 'Finance', 'normal': 'normal'}

_AR_DIGITS = str.maketrans('٠١٢٣٤٥٦٧٨٩۰۱۲۳۴۵۶۷۸۹', '01234567890123456789')


def _hijri_to_gregorian(y, m, d):
    """Tabular Islamic calendar → Gregorian (within a day of Umm al-Qura; the sheet used ar-SA dates)."""
    jd = (11 * y + 3) // 30 + 354 * y + 30 * m - (m - 1) // 2 + d + 1948440 - 385
    a = jd + 32044
    b = (4 * a + 3) // 146097
    c = a - 146097 * b // 4
    dd = (4 * c + 3) // 1461
    e = c - 1461 * dd // 4
    mm = (5 * e + 2) // 153
    return date(100 * b + dd - 4800 + mm // 10, mm + 3 - 12 * (mm // 10), e - (153 * mm + 2) // 5 + 1)


def parse_any_date(v):
    """Excel dates, ISO, d/m/Y, and the script's ar-SA strings such as «٢٠‏/٤‏/١٤٤٧ هـ»."""
    if v in (None, ''):
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).translate(_AR_DIGITS)
    nums = [int(n) for n in re.findall(r'\d+', s)]
    if len(nums) < 3:
        return None
    try:
        if nums[0] > 31:                       # Y-M-D
            y, m, d = nums[0], nums[1], nums[2]
        else:                                  # D/M/Y
            d, m, y = nums[0], nums[1], nums[2]
        if 1300 <= y < 1600:
            return _hijri_to_gregorian(y, m, d)
        return date(y, m, d)
    except ValueError:
        return None


def _sheet_rows(wb, name):
    for ws in wb.worksheets:
        if norm(ws.title) == norm(name):
            return [[c for c in r] for r in ws.iter_rows(values_only=True) if any(x not in (None, '') for x in r)]
    return []


def _records(rows, fields):
    clean = [['' if c is None else (c if isinstance(c, (date, datetime)) else str(c).strip()) for c in r] for r in rows]
    # to_records works on strings; keep the date objects by their position
    as_text = [[c.isoformat() if isinstance(c, (date, datetime)) else c for c in r] for r in clean]
    recs, _ = to_records(as_text, fields)
    return recs


def import_workbook(data):
    """``data``: bytes of the .xlsx. Must run inside the organization's tenant_context. → report dict."""
    from openpyxl import load_workbook
    from blueprints.custody import import_records, FIELDS as CUSTODY_FIELDS
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    rep = {'departments': 0, 'forms': 0, 'responses': 0, 'employees': 0, 'custody': 0, 'skipped': 0, 'problems': []}

    # departments
    by_legacy = {}
    for r in _records(_sheet_rows(wb, 'الاقسام'), DEPT_FIELDS):
        if not r.get('name'):
            continue
        code = (r.get('code') or r.get('legacy_id') or r['name'])[:50]
        d = Department.query.filter_by(code=code).first()
        if d is None:
            d = Department(code=code, name=r['name'][:200], icon=(r.get('icon') or '🏢')[:20],
                           manager_name=r.get('manager') or None, emails=r.get('email') or None,
                           dept_type=DEPT_TYPE.get((r.get('type') or 'normal').strip().lower(), 'normal'),
                           legacy_id=r.get('legacy_id') or None, is_active=True)
            db.session.add(d)
            db.session.flush()
            rep['departments'] += 1
        elif r.get('legacy_id') and not d.legacy_id:
            d.legacy_id = r['legacy_id']
        for key in (r.get('legacy_id'), r.get('code'), r.get('name')):
            if key:
                by_legacy[str(key)] = d
    for d in Department.query.all():
        for key in (d.legacy_id, d.code, d.name):
            if key:
                by_legacy.setdefault(str(key), d)

    # responses grouped by form
    responses = {}
    for r in _records(_sheet_rows(wb, 'الردود'), RESP_FIELDS):
        if r.get('form_id'):
            responses.setdefault(r['form_id'], []).append(r)

    # forms (+ roster)
    for r in _records(_sheet_rows(wb, 'النماذج'), FORM_FIELDS):
        if not r.get('form_id') or not r.get('name'):
            continue
        if ClearanceForm.query.filter_by(legacy_id=r['form_id']).first():
            rep['skipped'] += 1
            continue
        emp = None
        if r.get('number'):
            emp = Employee.query.filter_by(employee_number=str(r['number'])).first()
        if emp is None and r.get('email'):
            emp = Employee.query.filter(db.func.lower(Employee.email) == r['email'].lower()).first()
        if emp is None:
            emp = Employee(full_name=r['name'][:200], employee_number=str(r['number'])[:50] if r.get('number') else None,
                           email=r.get('email') or None, is_active=False, notes='مستورد من نماذج براءة الذمة القديمة')
            dept = by_legacy.get(r.get('department') or '')
            if dept is None and r.get('department'):
                dept = next((d for d in Department.query.all() if norm(d.name) == norm(r['department'])), None)
            emp.department_id = dept.id if dept else None
            db.session.add(emp)
            db.session.flush()
            rep['employees'] += 1
        status = FORM_STATUS.get(norm(r.get('status')), 'in_progress')
        note = r.get('hr_note') or ''
        settled = status == 'closed' and note.startswith('تسوية محاسبة')
        created = parse_any_date(r.get('created'))
        ref = r.get('ref') or f'CLR-OLD-{r["form_id"][-6:]}'
        if ClearanceForm.query.filter_by(ref_code=ref).first():
            ref = f'{ref}-{r["form_id"][-4:]}'
        form = ClearanceForm(legacy_id=r['form_id'], ref_code=ref[:50], employee_id=emp.id, employee_name=r['name'][:200],
                             employee_number=str(r['number'])[:50] if r.get('number') else None,
                             employee_email=r.get('email') or None, department_name=r.get('department') or None,
                             resign_date=parse_any_date(r.get('resign')), status=status,
                             hr_note=None if settled else (note or None),
                             finance_note=note.replace('تسوية محاسبة:', '').strip() or None if settled else None,
                             created_at=datetime.combine(created, datetime.min.time()) if created else utcnow(),
                             created_by='استيراد من Google Sheet')
        if settled:
            form.settled_at = form.closed_at = form.created_at
            form.settled_by = 'المحاسبة'
        elif status == 'closed':
            form.closed_at = form.created_at
        db.session.add(form)
        db.session.flush()
        seen = set()
        for rr in responses.get(r['form_id'], []):
            d = by_legacy.get(str(rr.get('dept_id') or ''))
            if d is None or d.id in seen:
                if d is None:
                    rep['problems'].append(f'{ref}: قسم غير معروف «{rr.get("dept_id")}»')
                continue
            seen.add(d.id)
            at = parse_any_date(rr.get('at'))
            db.session.add(ClearanceResponse(form_id=form.id, department_id=d.id,
                                             status=RESP_STATUS.get(norm(rr.get('status')), 'pending'),
                                             note=(rr.get('note') or None), responded_by=rr.get('by') or None,
                                             responded_at=datetime.combine(at, datetime.min.time()) if at else None))
            rep['responses'] += 1
        rep['forms'] += 1
    db.session.commit()

    # custody (same importer as the web page; no e-mails)
    rows = _sheet_rows(wb, 'العهد')
    if rows:
        recs = _records(rows, CUSTODY_FIELDS)
        for rec in recs:   # the sheet keeps the department code in «كود_القسم»; dates may be ar-SA strings
            for k in ('handed_at', 'returned_at'):
                if rec.get(k):
                    dd = parse_any_date(rec[k])
                    rec[k] = dd.isoformat() if dd else ''
        res = import_records(recs, None)
        rep['custody'] = res['added']
        rep['skipped'] += res['skipped']
        rep['problems'] += res['problems']
    wb.close()
    return rep

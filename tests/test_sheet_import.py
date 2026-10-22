"""Import of the old Google Sheet workbook (as downloaded from Google Sheets)."""
import io
from datetime import date

from openpyxl import Workbook

from extensions import db
from models import Department, ClearanceForm, Custody, Employee
from tenancy import tenant_context
from tests.conftest import login
from services.sheets_import import parse_any_date


def _workbook():
    wb = Workbook()
    ws = wb.active
    ws.title = 'الاقسام'
    ws.append(['id', 'code', 'name', 'icon', 'manager', 'email', 'type'])
    ws.append(['D1', 'IT9', 'التقنية', '💻', 'منى', 'it@old.test', 'normal'])
    ws.append(['D2', 'HR9', 'شؤون الموظفين', '👥', 'سارة', 'hr@old.test', 'HR'])
    ws = wb.create_sheet('النماذج')
    ws.append(['formId', 'formCode', 'empCode', 'employeeName', 'employeeId', 'employeeEmail', 'department',
               'resignDate', 'createdAt', 'status', 'hrNote'])
    ws.append(['FORM-AAA', 'CLR-2025-AAA', 'EMP-1', 'قديم أحمد', '55', 'old@a.test', 'التقنية', date(2025, 5, 1),
               '٢٠‏/٤‏/١٤٤٧ هـ', 'مقفل', 'تسوية محاسبة: خُصمت القيمة'])
    ws = wb.create_sheet('الردود')
    ws.append(['formId', 'deptId', 'status', 'note', 'respondedAt', 'respondedBy'])
    ws.append(['FORM-AAA', 'D1', 'معلق', 'لم يُرجع الحاسوب', '', 'منى'])
    ws.append(['FORM-AAA', 'D2', 'خالية الذمة', '', '', 'سارة'])
    ws = wb.create_sheet('العهد')
    ws.append(['معرف_العهدة', 'العهدة', 'المستلم', 'بريد_المستلم', 'القسم', 'كود_القسم', 'مسلم_العهدة',
               'تاريخ_التسليم', 'الحالة', 'تاريخ_الإرجاع', 'مستلم_الإرجاع', 'بريد_مسلم_العهدة'])
    ws.append(['CUST-OLD1', 'حاسوب', 'قديم أحمد', 'old@a.test', 'التقنية', 'IT9', 'منى', '١/٣/١٤٤٦ هـ', 'مسلمة', '', '', ''])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_hijri_dates():
    assert parse_any_date('٢٠‏/٤‏/١٤٤٧ هـ') in (date(2025, 10, 11), date(2025, 10, 12), date(2025, 10, 13))
    assert parse_any_date('2026-01-05') == date(2026, 1, 5)
    assert parse_any_date('5/1/2026') == date(2026, 1, 5)


def test_import_workbook_is_idempotent(app, two_tenants):
    t = two_tenants
    c = app.test_client()
    login(c, 'admin', org='a')
    data = _workbook()
    for _ in range(2):
        r = c.post('/admin/settings/import-sheet', data={'file': (io.BytesIO(data), 'old.xlsx')},
                   content_type='multipart/form-data', follow_redirects=True)
        assert r.status_code == 200
    with tenant_context(t['a']):
        assert Department.query.filter_by(code='HR9').one().dept_type == 'HR'
        f = ClearanceForm.query.one()
        assert f.status == 'closed' and f.settled and f.finance_note == 'خُصمت القيمة'
        assert sorted(r.status for r in f.responses) == ['cleared', 'held']
        emp = Employee.query.filter_by(employee_number='55').one()
        assert not emp.is_active and f.employee_id == emp.id
        cu = Custody.query.one()
        assert cu.employee_id == emp.id and cu.department.code == 'IT9' and cu.handed_at.year == 2024
    with tenant_context(t['b']):
        assert ClearanceForm.query.count() == 0
    assert app.config['MAIL_OUTBOX'] == []
